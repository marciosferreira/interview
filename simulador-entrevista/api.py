import asyncio
import functools
import hashlib
import mimetypes
import os
import struct
import time
import traceback
import uuid
from pathlib import Path
from typing import Optional

from dotenv import load_dotenv
load_dotenv()

import stripe as _stripe

from fastapi import BackgroundTasks, FastAPI, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect, Depends
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.types import ASGIApp, Receive, Scope, Send
from fastapi.responses import Response, FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from langgraph.types import Command
from openai import AsyncOpenAI
from pydantic import BaseModel

from simulator import (
    graph, make_initial_state, model, get_model, session_store,
    _build_report_system, REPORT_PHASE_MAP, get_db_conn,
    _EXPLORER_MODEL, _HUNTER_MODEL, extract_hire_signal,
)
from auth import (
    UserCreate, UserLogin, TokenResponse, RegisterResponse, ProfileUpdate,
    VERIFICATION_TOKEN_EXPIRES_HOURS,
    create_user, authenticate_user, get_user_by_id, update_profile,
    create_verification_token, verify_email_token,
    create_reset_token, reset_password_with_token,
    upgrade_plan, count_interviews_this_week, PLAN_WEEKLY_LIMITS,
    create_job_session, update_job_session_context, get_job_session, list_job_sessions,
    reset_weekly_limit,
    create_access_token, get_current_user, decode_token_raw,
    create_contact_message, list_contact_messages, save_contact_reply, get_contact_message, delete_contact_message,
    set_stripe_info, clear_stripe_subscription, get_user_by_stripe_customer, get_user_by_email,
    get_stripe_info, set_stripe_cancel_at, force_verify_email, delete_account,
)
from email_service import send_verification_email, send_reset_email, send_contact_notification, send_contact_reply
from prompt_generator import generate_interview_context, extract_candidate_name

# DB helper — works for both psycopg2 pool (Postgres) and sqlite3
_is_pg = bool(os.getenv("DATABASE_URL"))
_ph = "%s" if _is_pg else "?"


def _with_conn(func):
    """Borrow a DB connection, run func(conn), return it to the pool."""
    with get_db_conn() as conn:
        return func(conn)


def _db_exec(sql: str, params=(), conn=None):
    """Execute SQL; conn must be provided for postgres (use within get_db_conn context)."""
    if conn is not None:
        if _is_pg:
            cur = conn.cursor()
            cur.execute(sql, params)
            return cur
        return conn.execute(sql, params)
    # SQLite fallback (conn is the shared sqlite connection yielded by get_db_conn)
    with get_db_conn() as c:
        return c.execute(sql, params)

_openai_api_key = os.getenv("OPENAI_API_KEY")
openai_client = AsyncOpenAI(api_key=_openai_api_key) if _openai_api_key else None

from google import genai as _genai
from google.genai import types as _genai_types
_gemini_client = _genai.Client(api_key=os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY"))

_TTS_PROVIDER = os.getenv("TTS_PROVIDER", "openai").lower()
_GEMINI_TTS_VOICE = os.getenv("GEMINI_TTS_VOICE", "Zephyr")
_GEMINI_TTS_SCORECARD_VOICE = os.getenv("GEMINI_TTS_SCORECARD_VOICE", _GEMINI_TTS_VOICE)

_stripe.api_key = os.getenv("STRIPE_SECRET_KEY", "")
_STRIPE_WEBHOOK_SECRET = os.getenv("STRIPE_WEBHOOK_SECRET", "")
_STRIPE_PRICE_ID_USD   = os.getenv("STRIPE_PRICE_ID_USD", os.getenv("STRIPE_PRICE_ID", ""))
_STRIPE_PRICE_ID_BRL   = os.getenv("STRIPE_PRICE_ID_BRL", "")
_APP_BASE_URL          = os.getenv("APP_BASE_URL", "http://localhost:8001").rstrip("/")

_COUNTRY_NAMES = {
    "US": "United States",
    "BR": "Brazil",
    "CA": "Canada",
    "GB": "United Kingdom",
    "PT": "Portugal",
    "ES": "Spain",
    "FR": "France",
    "DE": "Germany",
    "IT": "Italy",
    "NL": "Netherlands",
    "IE": "Ireland",
    "AU": "Australia",
    "IN": "India",
    "MX": "Mexico",
    "AR": "Argentina",
    "CL": "Chile",
    "CO": "Colombia",
    "PE": "Peru",
}


def _stripe_price_for(language: str) -> str:
    if language == "pt" and _STRIPE_PRICE_ID_BRL:
        return _STRIPE_PRICE_ID_BRL
    return _STRIPE_PRICE_ID_USD


def _client_ip(request: Request) -> str:
    forwarded_for = request.headers.get("x-forwarded-for", "")
    if forwarded_for:
        return forwarded_for.split(",", 1)[0].strip()
    real_ip = request.headers.get("x-real-ip") or request.headers.get("cf-connecting-ip")
    if real_ip:
        return real_ip.strip()
    return request.client.host if request.client else ""


def _signup_country(request: Request) -> tuple[str, str]:
    code = (
        request.headers.get("cloudfront-viewer-country") or
        request.headers.get("cf-ipcountry") or
        request.headers.get("x-vercel-ip-country") or
        ""
    ).strip().upper()
    if len(code) != 2 or code == "XX":
        return "", ""
    return code, _COUNTRY_NAMES.get(code, code)


app = FastAPI()


def _parse_cookie_header(cookie_header: str) -> dict[str, str]:
    cookies = {}
    for item in cookie_header.split(";"):
        if "=" not in item:
            continue
        key, value = item.split("=", 1)
        cookies[key.strip()] = value.strip()
    return cookies


def _detect_supported_language(accept_language: str, fallback: str = "en") -> str:
    accepted = []
    for index, item in enumerate(accept_language.split(",")):
        parts = item.strip().split(";")
        if not parts or not parts[0]:
            continue
        lang = parts[0].split("-")[0].lower()
        if lang not in ("en", "pt"):
            continue
        quality = 1.0
        for part in parts[1:]:
            part = part.strip()
            if part.startswith("q="):
                try:
                    quality = float(part[2:])
                except ValueError:
                    quality = 0.0
        accepted.append((quality, -index, lang))

    if accepted:
        accepted.sort(reverse=True)
        return accepted[0][2]
    return fallback


def _preferred_landing_language_from_headers(headers: dict[str, str]) -> str:
    cookies = _parse_cookie_header(headers.get("cookie", ""))
    cookie_lang = cookies.get("preferred_lang")
    if cookie_lang in ("en", "pt"):
        return cookie_lang
    return _detect_supported_language(headers.get("accept-language", ""))


def _preferred_landing_language(request: Request) -> str:
    return _preferred_landing_language_from_headers(dict(request.headers))


def _sync_user_language(conn, user: dict, language: Optional[str]) -> dict:
    if language not in ("en", "pt") or user.get("language") == language:
        return user

    cur = conn.cursor()
    cur.execute(f"UPDATE users SET language = {_ph} WHERE id = {_ph}", (language, user["id"]))
    cur.close()
    conn.commit()
    return {**user, "language": language}


class _CacheMiddleware:
    def __init__(self, app: ASGIApp) -> None:
        self.app = app

    async def __call__(self, scope: Scope, receive: Receive, send: Send) -> None:
        if scope["type"] != "http":
            print(f"[middleware] passing through scope type: {scope['type']}")
            await self.app(scope, receive, send)
            return

        path = scope.get("path", "")

        headers = {
            key.decode("latin-1").lower(): value.decode("latin-1")
            for key, value in scope.get("headers", [])
        }

        # The public root is only a locale negotiator. Canonical landing URLs are
        # /en/ and /pt/ so crawlers, caches, and users all see a stable URL.
        if path in ("", "/"):
            lang = _preferred_landing_language_from_headers(headers)
            response = RedirectResponse(url=f"/{lang}/", status_code=302)
            response.headers["Vary"] = "Accept-Language, Cookie"
            await response(scope, receive, send)
            return

        # Serve index.html for language landing routes before StaticFiles intercepts
        if path in ("/en", "/en/", "/pt", "/pt/"):
            index_path = Path(__file__).parent / "static" / "index.html"
            response = FileResponse(str(index_path))
            await response(scope, receive, send)
            return

        no_cache = path.endswith(".html") or path in ("/", "", "/en", "/en/", "/pt", "/pt/") or path.endswith(".json")
        cacheable_static = path.lower().endswith((
            ".css", ".js", ".png", ".jpg", ".jpeg", ".webp", ".avif",
            ".svg", ".ico", ".woff2", ".wav", ".mp3",
        ))

        async def send_with_headers(message):
            if message["type"] == "http.response.start":
                headers = list(message.get("headers", []))
                if no_cache:
                    headers += [
                        (b"cache-control", b"no-cache, no-store, must-revalidate"),
                        (b"pragma", b"no-cache"),
                    ]
                elif cacheable_static:
                    headers += [
                        (b"cache-control", b"public, max-age=31536000, immutable"),
                    ]
                message = {**message, "headers": headers}
            await send(message)

        await self.app(scope, receive, send_with_headers)

app.add_middleware(_CacheMiddleware)


# ── Auth endpoints ────────────────────────────────────────────────────────────

@app.post("/auth/register", response_model=RegisterResponse)
async def register(body: UserCreate, background_tasks: BackgroundTasks, request: Request):
    honeypot_filled = (
        (body.last_name and body.last_name.strip()) or
        (body.company_site and body.company_site.strip()) or
        (body.extra_context and body.extra_context.strip())
    )
    if honeypot_filled:
        return RegisterResponse(email=body.email.strip().lower())
    if not body.name.strip():
        raise HTTPException(status_code=422, detail="Name is required")
    if not body.email.strip() or "@" not in body.email:
        raise HTTPException(status_code=422, detail="Valid email is required")
    if len(body.password) < 8:
        raise HTTPException(status_code=422, detail="Password must be at least 8 characters")
    language = body.language if body.language in ("en", "pt") else "en"
    signup_ip = _client_ip(request)
    signup_country_code, signup_country_name = _signup_country(request)
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: create_user(c, body.name.strip(), body.email.strip(),
                                                  body.password, language,
                                                  signup_ip, signup_country_code,
                                                  signup_country_name)),
    )
    token = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: create_verification_token(c, user["id"]))
    )
    background_tasks.add_task(
        send_verification_email, user["email"], token, user.get("language", "en")
    )
    return RegisterResponse(email=user["email"])


@app.post("/auth/login", response_model=TokenResponse)
async def login(body: UserLogin):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: authenticate_user(c, body.email, body.password)),
    )
    if not user:
        raise HTTPException(status_code=401, detail="Invalid email or password")
    user = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: _sync_user_language(c, user, body.language)),
    )
    if not user.get("email_verified") and user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="email_not_verified")
    token = create_access_token(user["id"], user["email"])
    return TokenResponse(access_token=token, user=user)


@app.delete("/auth/account")
async def delete_my_account(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    stripe_info = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_stripe_info(c, current_user["id"])))
    is_hunter = user.get("plan") == "hunter"
    cancel_at = stripe_info.get("stripe_cancel_at") if stripe_info else None
    if is_hunter and not cancel_at:
        raise HTTPException(
            status_code=403,
            detail="You have an active Hunter subscription. Please cancel it first before deleting your account.",
        )
    if user.get("email") == ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="The admin account cannot be deleted.")
    await loop.run_in_executor(None, lambda: _with_conn(lambda c: delete_account(c, current_user["id"])))
    return {"ok": True}


@app.get("/auth/me")
async def me(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"]))
    )
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


@app.put("/auth/profile")
async def update_user_profile(body: ProfileUpdate, current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: update_profile(
            c, current_user["id"],
            body.name, body.email, body.current_password, body.new_password,
            body.language,
        )),
    )
    # If email changed, send new verification email
    if result.get("email_changed"):
        token = await loop.run_in_executor(
            None, lambda: _with_conn(lambda c: create_verification_token(c, current_user["id"]))
        )
        try:
            await loop.run_in_executor(
                None, lambda: send_verification_email(result["email"], token, result.get("language", "en"))
            )
        except Exception as exc:
            print(f"[profile] Failed to send verification email: {exc}")
    return result


@app.post("/auth/send-verification")
async def send_verification(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if user.get("email_verified"):
        return {"ok": True, "already_verified": True}
    token = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: create_verification_token(c, current_user["id"]))
    )
    mock = os.getenv("MOCK_EMAIL", "false").lower() == "true"
    try:
        await loop.run_in_executor(None, lambda: send_verification_email(user["email"], token, user.get("language", "en")))
    except Exception as exc:
        print(f"[send-verification] Failed: {exc}")
        raise HTTPException(status_code=503, detail="Failed to send email — please try again")
    response = {"ok": True}
    if mock:
        response["dev_token"] = token  # expose token in dev mode for manual testing
    return response


class ResendVerificationPublicRequest(BaseModel):
    email: str


_RESEND_COOLDOWN_SECONDS = 300  # 5 minutes between resend attempts


@app.post("/auth/resend-verification-public")
async def resend_verification_public(body: ResendVerificationPublicRequest):
    """Public endpoint — resend verification email without auth (used on register success screen)."""
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_email(c, body.email.strip().lower())))
    if not user or user.get("email_verified"):
        return {"ok": True}
    # Rate limit: refuse if last token was issued less than _RESEND_COOLDOWN_SECONDS ago
    expires = user.get("verification_expires")
    if expires:
        now_ms = int(time.time() * 1000)
        token_age_ms = VERIFICATION_TOKEN_EXPIRES_HOURS * 3_600_000 - (int(expires) - now_ms)
        if token_age_ms < _RESEND_COOLDOWN_SECONDS * 1000:
            wait_s = max(1, -(-(_RESEND_COOLDOWN_SECONDS * 1000 - token_age_ms) // 1000))  # ceiling div
            raise HTTPException(status_code=429, detail={"message": f"Please wait {wait_s} seconds before requesting another email.", "wait_seconds": wait_s})
    token = await loop.run_in_executor(None, lambda: _with_conn(lambda c: create_verification_token(c, user["id"])))
    try:
        await loop.run_in_executor(None, lambda: send_verification_email(user["email"], token, user.get("language", "en")))
    except Exception as exc:
        print(f"[resend-verification-public] Failed: {exc}")
        raise HTTPException(status_code=503, detail="Failed to send email — please try again")
    return {"ok": True}


@app.get("/auth/verify-email")
async def verify_email(token: str):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(None, lambda: _with_conn(lambda c: verify_email_token(c, token)))
    if not result:
        raise HTTPException(status_code=400, detail="Invalid or expired verification link")
    return {"ok": True, "email": result["email"]}


class ForgotPasswordRequest(BaseModel):
    email: str


class ResetPasswordRequest(BaseModel):
    token: str
    new_password: str


@app.post("/auth/forgot-password")
async def forgot_password(body: ForgotPasswordRequest):
    loop = asyncio.get_running_loop()
    result = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: create_reset_token(c, body.email))
    )
    mock = os.getenv("MOCK_EMAIL", "false").lower() == "true"
    response = {"ok": True}  # always return ok to avoid email enumeration
    if result:
        _, token, lang = result
        try:
            await loop.run_in_executor(None, lambda: send_reset_email(body.email, token, lang))
        except Exception as exc:
            print(f"[forgot-password] Failed to send email: {exc}")
        if mock:
            response["dev_token"] = token
    return response


@app.post("/auth/reset-password")
async def reset_password(body: ResetPasswordRequest):
    loop = asyncio.get_running_loop()
    ok = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: reset_password_with_token(c, body.token, body.new_password))
    )
    if not ok:
        raise HTTPException(status_code=400, detail="Invalid or expired reset link")
    return {"ok": True}


@app.post("/auth/upgrade")
async def upgrade_to_hunter(current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Use Stripe checkout to upgrade")
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: _with_conn(lambda c: upgrade_plan(c, current_user["id"], "hunter")))
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    return {"ok": True, "plan": "hunter", "user": user}


@app.get("/auth/usage")
async def get_usage(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    plan  = user.get("plan", "free") if user else "free"
    used  = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: count_interviews_this_week(c, current_user["id"]))
    )
    limit = PLAN_WEEKLY_LIMITS.get(plan, 3)
    stripe_info = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_stripe_info(c, current_user["id"])))
    return {
        "plan": plan, "used": used, "limit": limit, "remaining": max(0, limit - used),
        "has_stripe_subscription": bool(stripe_info.get("stripe_subscription_id")),
        "stripe_cancel_at": stripe_info.get("stripe_cancel_at"),
    }


@app.post("/auth/downgrade")
async def downgrade_to_free(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: _with_conn(lambda c: upgrade_plan(c, current_user["id"], "free")))
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    return {"ok": True, "plan": "free", "user": user}


# ── Stripe endpoints ──────────────────────────────────────────────────────────

@app.post("/stripe/create-checkout-session")
async def stripe_create_checkout(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"])))
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    price_id = _stripe_price_for(user.get("language", "en"))
    if not price_id:
        raise HTTPException(status_code=503, detail="Stripe not configured")
    if user.get("plan") == "hunter":
        raise HTTPException(status_code=400, detail="Already on Hunter plan")
    try:
        session = _stripe.checkout.Session.create(
            payment_method_types=["card"],
            mode="subscription",
            line_items=[{"price": price_id, "quantity": 1}],
            success_url=f"{_APP_BASE_URL}/upgrade-success.html?session_id={{CHECKOUT_SESSION_ID}}",
            cancel_url=f"{_APP_BASE_URL}/upgrade.html",
            client_reference_id=user["id"],
            customer_email=user.get("email"),
            metadata={"user_id": user["id"]},
        )
        return {"url": session.url}
    except _stripe.error.StripeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))


@app.post("/stripe/webhook")
async def stripe_webhook(request: Request):
    payload    = await request.body()
    sig_header = request.headers.get("stripe-signature", "")
    try:
        event = _stripe.Webhook.construct_event(payload, sig_header, _STRIPE_WEBHOOK_SECRET)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid payload")
    except _stripe.error.SignatureVerificationError:
        raise HTTPException(status_code=400, detail="Invalid signature")

    try:
        etype = event["type"]
        data  = event["data"]["object"]
        print(f"[Stripe] received event: {etype} id={event.id}")
        _handle_stripe_event(etype, data)
    except Exception:
        tb = traceback.format_exc()
        print(f"[Stripe] UNHANDLED ERROR:\n{tb}")
        return {"ok": False, "error": tb}

    return {"ok": True}


def _find_user_for_customer(customer_id: str) -> Optional[dict]:
    """Look up local user by Stripe customer ID, falling back to customer email."""
    user = _with_conn(lambda c: get_user_by_stripe_customer(c, customer_id))
    if user:
        return user
    # Fallback: fetch customer from Stripe and match by email.
    try:
        cus = _stripe.Customer.retrieve(customer_id)
        email = _sg(cus, "email") or ""
        if email:
            user = _with_conn(lambda c: get_user_by_email(c, email))
            if user:
                print(f"[Stripe] found user {user['id']} via email fallback for customer {customer_id}")
    except Exception as exc:
        print(f"[Stripe] could not fetch customer {customer_id}: {exc}")
    return user


def _sg(obj, key, default=None):
    """Safe getter for Stripe SDK v15 typed objects that lack .get()."""
    try:
        return getattr(obj, key, default)
    except Exception:
        return default


def _handle_stripe_event(etype: str, data) -> None:
    if etype == "checkout.session.completed":
        meta            = _sg(data, "metadata") or {}
        user_id         = _sg(data, "client_reference_id") or meta.get("user_id")
        customer_id     = _sg(data, "customer")
        subscription_id = _sg(data, "subscription")
        print(f"[Stripe] checkout.session.completed user_id={user_id} customer={customer_id} sub={subscription_id}")
        if user_id and customer_id:
            _with_conn(lambda c: upgrade_plan(c, user_id, "hunter"))
            _with_conn(lambda c: set_stripe_info(c, user_id, customer_id, subscription_id or ""))
            print(f"[Stripe] → upgraded user {user_id}")
        else:
            print(f"[Stripe] → missing user_id or customer_id, cannot upgrade")

    elif etype in ("invoice_payment.paid", "invoice.paid"):
        # invoice_payment.paid (new API) has no customer field — must fetch invoice.
        invoice_id      = _sg(data, "invoice") or _sg(data, "id")
        customer_id     = _sg(data, "customer")
        subscription_id = _sg(data, "subscription")
        print(f"[Stripe] {etype} invoice={invoice_id} customer={customer_id}")

        # Resolve customer_id and subscription_id from the invoice if missing.
        if invoice_id and (not customer_id or not subscription_id):
            try:
                inv             = _stripe.Invoice.retrieve(invoice_id)
                customer_id     = customer_id or _sg(inv, "customer")
                subscription_id = subscription_id or _sg(inv, "subscription")
                print(f"[Stripe] fetched invoice → customer={customer_id} sub={subscription_id}")
            except Exception as exc:
                print(f"[Stripe] could not fetch invoice {invoice_id}: {exc}")

        if not customer_id:
            print(f"[Stripe] {etype} → no customer_id resolved, skipping")
            return

        user = _find_user_for_customer(customer_id)
        if not user:
            print(f"[Stripe] {etype} → no local user for customer {customer_id}")
            return

        # Always ensure customer mapping is stored (may be missing if checkout event failed).
        if subscription_id:
            _with_conn(lambda c: set_stripe_info(c, user["id"], customer_id, subscription_id))

        if user.get("plan") != "hunter":
            _with_conn(lambda c: upgrade_plan(c, user["id"], "hunter"))
            print(f"[Stripe] {etype} → upgraded user {user['id']} to hunter")
        else:
            print(f"[Stripe] {etype} → user {user['id']} already hunter")

    elif etype == "customer.subscription.deleted":
        customer_id = _sg(data, "customer")
        user = _find_user_for_customer(customer_id) if customer_id else None
        if user:
            _with_conn(lambda c: upgrade_plan(c, user["id"], "free"))
            _with_conn(lambda c: clear_stripe_subscription(c, user["id"]))
            _with_conn(lambda c: set_stripe_cancel_at(c, user["id"], None))
            print(f"[Stripe] subscription.deleted → downgraded user {user['id']} to free")

    elif etype == "customer.subscription.updated":
        customer_id = _sg(data, "customer")
        status      = _sg(data, "status")
        user = _find_user_for_customer(customer_id) if customer_id else None
        if user and status not in ("active", "trialing"):
            _with_conn(lambda c: upgrade_plan(c, user["id"], "free"))
            print(f"[Stripe] subscription.updated status={status} → downgraded user {user['id']} to free")

    elif etype == "invoice.payment_failed":
        customer_id = _sg(data, "customer")
        attempt     = _sg(data, "attempt_count") or 1
        print(f"[Stripe] payment failed (attempt {attempt}) for customer {customer_id}")

    else:
        print(f"[Stripe] unhandled event type: {etype}")


@app.get("/stripe/invoices")
async def stripe_invoices(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    info = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_stripe_info(c, current_user["id"])))
    customer_id = info.get("stripe_customer_id")
    if not customer_id:
        return {"invoices": []}
    try:
        result = _stripe.Invoice.list(customer=customer_id, limit=24, status="paid")
        data_list = list(result.auto_paging_iter()) if hasattr(result, "auto_paging_iter") else (_sg(result, "data") or [])
        invoices = [
            {
                "id":                  _sg(inv, "id"),
                "created":             _sg(inv, "created"),
                "amount_paid":         _sg(inv, "amount_paid"),
                "currency":            _sg(inv, "currency"),
                "invoice_pdf":         _sg(inv, "invoice_pdf"),
                "hosted_invoice_url":  _sg(inv, "hosted_invoice_url"),
            }
            for inv in data_list
            if _sg(inv, "id")
        ]
        return {"invoices": invoices}
    except _stripe.error.StripeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))


@app.post("/stripe/cancel")
async def stripe_cancel_subscription(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    info = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_stripe_info(c, current_user["id"])))
    subscription_id = info.get("stripe_subscription_id")
    if not subscription_id:
        raise HTTPException(status_code=400, detail="No active Stripe subscription found")
    try:
        sub = _stripe.Subscription.modify(subscription_id, cancel_at_period_end=True)
        period_end = _sg(sub, "current_period_end")
        _with_conn(lambda c: set_stripe_cancel_at(c, current_user["id"], period_end))
        return {"ok": True, "period_end": period_end}
    except _stripe.error.StripeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))


@app.post("/stripe/portal")
async def stripe_portal(current_user: dict = Depends(get_current_user)):
    if not _STRIPE_PRICE_ID:
        raise HTTPException(status_code=503, detail="Stripe not configured")
    loop = asyncio.get_running_loop()
    info = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_stripe_info(c, current_user["id"])))
    customer_id = info.get("stripe_customer_id")
    if not customer_id:
        raise HTTPException(status_code=400, detail="No Stripe subscription found")
    try:
        portal = _stripe.billing_portal.Session.create(
            customer=customer_id,
            return_url=f"{_APP_BASE_URL}/upgrade.html",
        )
        return {"url": portal.url}
    except _stripe.error.StripeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))


# ── Job session / prepare endpoints ──────────────────────────────────────────

class PrepareRequest(BaseModel):
    job_title:       str
    company:         str = ""
    job_description: str
    resume_text:     str


@app.post("/prepare")
async def prepare(body: PrepareRequest, current_user: dict = Depends(get_current_user)):
    """Generate a personalized interview context and create a job session."""
    if len(body.job_description.strip()) < 50:
        raise HTTPException(status_code=422, detail="Job description is too short")
    if len(body.resume_text.strip()) < 50:
        raise HTTPException(status_code=422, detail="Resume text is too short")

    loop = asyncio.get_running_loop()

    # Fetch user info (language + plan)
    user = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"]))
    )
    if user and not user.get("email_verified") and current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="email_not_verified")
    language = user.get("language", "en") if user else "en"
    plan     = user.get("plan", "free")   if user else "free"

    # Enforce weekly interview limit
    week_count = await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: count_interviews_this_week(c, current_user["id"]))
    )
    limit = PLAN_WEEKLY_LIMITS.get(plan, 3)
    if week_count >= limit:
        raise HTTPException(
            status_code=429,
            detail={
                "code": "weekly_limit_reached",
                "plan": plan,
                "limit": limit,
                "used": week_count,
                "language": language,
                "message": f"You've used all {limit} interviews this week on your {plan.capitalize()} plan.",
            },
        )

    # Generate personalized interview context (LLM call — may take 10–20 seconds)
    try:
        interview_context = await generate_interview_context(
            job_title=body.job_title,
            company=body.company,
            job_description=body.job_description,
            resume_text=body.resume_text,
            language=language,
            model_name=_HUNTER_MODEL if plan == "hunter" else _EXPLORER_MODEL,
        )
    except Exception as exc:
        print(f"[prepare] LLM generation failed: {exc}")
        raise HTTPException(
            status_code=503,
            detail="Failed to generate interview context — please try again in a moment.",
        )

    candidate_name = extract_candidate_name(interview_context)

    # Persist in DB
    job_session_id = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: create_job_session(
            c,
            user_id=current_user["id"],
            job_title=body.job_title,
            company=body.company,
            job_description=body.job_description,
            resume_text=body.resume_text,
            interview_context=interview_context,
        )),
    )

    return {
        "job_session_id": job_session_id,
        "candidate_name": candidate_name,
        "job_title": body.job_title,
        "company": body.company,
        "context_preview": interview_context[:500] + "…" if len(interview_context) > 500 else interview_context,
    }


@app.get("/job-sessions")
async def list_user_job_sessions(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: list_job_sessions(c, current_user["id"]))
    )


# ── Session history endpoints ─────────────────────────────────────────────────

class SessionUpsert(BaseModel):
    startedAt:    int
    lastActiveAt: int
    fase:         str
    history:      list
    jobSessionId: str = ""


@app.get("/sessions")
async def list_sessions(current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(
        None, lambda: session_store.list_sessions(user_id=current_user["id"])
    )


@app.put("/sessions/{thread_id}")
async def upsert_session(thread_id: str, body: SessionUpsert,
                         current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(
        None,
        lambda: session_store.upsert(
            thread_id,
            body.startedAt,
            body.lastActiveAt,
            body.fase,
            body.history,
            user_id=current_user["id"],
            job_session_id=body.jobSessionId or None,
        ),
    )
    return {"ok": True}


@app.delete("/sessions/{thread_id}")
async def delete_session(thread_id: str, current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(
        None, lambda: session_store.delete(thread_id, user_id=current_user["id"])
    )
    return {"ok": True}


ADMIN_EMAIL = "marciosferreira@yahoo.com.br"


@app.get("/admin/users")
async def admin_list_users(current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    loop = asyncio.get_running_loop()
    week_ago_ms = int((__import__("time").time() - 7 * 24 * 3600) * 1000)

    def _query():
        with get_db_conn() as conn:
            rows = _db_exec(f"""
                SELECT
                    u.id, u.name, u.email, u.plan, u.email_verified, u.created_at,
                    u.signup_ip, u.signup_country_code, u.signup_country_name,
                    COUNT(sm.thread_id)                                          AS total_sessions,
                    SUM(CASE WHEN sm.fase = 'done' THEN 1 ELSE 0 END)           AS completed,
                    SUM(CASE WHEN sm.started_at >= CASE
                            WHEN COALESCE(u.week_reset_at, 0) > {week_ago_ms}
                            THEN u.week_reset_at ELSE {week_ago_ms}
                        END THEN 1 ELSE 0 END)                                  AS sessions_week,
                    MAX(sm.last_active_at)                                       AS last_active,
                    u.stripe_subscription_id
                FROM users u
                LEFT JOIN session_meta sm ON u.id = sm.user_id
                GROUP BY u.id, u.stripe_subscription_id
                ORDER BY u.created_at DESC
            """, (), conn=conn).fetchall()
        return [
            {
                "id": r[0], "name": r[1], "email": r[2],
                "plan": r[3] or "free", "email_verified": bool(r[4]),
                "created_at": r[5],
                "signup_ip": r[6] or "", "signup_country_code": r[7] or "",
                "signup_country_name": r[8] or "",
                "total_sessions": r[9] or 0, "completed": r[10] or 0,
                "sessions_week": r[11] or 0, "last_active": r[12],
                "stripe_subscription_id": r[13] or "",
            }
            for r in rows
        ]

    users = await loop.run_in_executor(None, _query)
    return {"users": users}


@app.post("/admin/users/{user_id}/plan")
async def admin_set_plan(user_id: str, body: dict, current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    plan = body.get("plan")
    if plan not in ("free", "hunter"):
        raise HTTPException(status_code=422, detail="Invalid plan")
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: _with_conn(lambda c: upgrade_plan(c, user_id, plan)))
    return {"ok": True}


@app.post("/admin/users/{user_id}/verify")
async def admin_verify_email(user_id: str, current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    loop = asyncio.get_running_loop()
    ok = await loop.run_in_executor(None, lambda: _with_conn(lambda c: force_verify_email(c, user_id)))
    if not ok:
        raise HTTPException(status_code=404, detail="User not found")
    return {"ok": True}


@app.delete("/admin/users/{user_id}")
async def admin_delete_user(user_id: str, current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")

    def _delete():
        with get_db_conn() as conn:
            # Guard: never delete the admin account
            row = _db_exec(f"SELECT email FROM users WHERE id = {_ph}", (user_id,), conn=conn).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="User not found")
            if row[0] == ADMIN_EMAIL:
                raise HTTPException(status_code=403, detail="Cannot delete the admin account")
            # Delete all interview data, then the user
            thread_ids = [r[0] for r in _db_exec(
                f"SELECT thread_id FROM session_meta WHERE user_id = {_ph}", (user_id,), conn=conn
            ).fetchall()]
            if thread_ids:
                placeholders = ",".join([_ph] * len(thread_ids))
                _db_exec(f"DELETE FROM session_meta WHERE thread_id IN ({placeholders})", tuple(thread_ids), conn=conn)
            _db_exec(f"DELETE FROM job_sessions WHERE user_id = {_ph}", (user_id,), conn=conn)
            _db_exec(f"DELETE FROM users WHERE id = {_ph}", (user_id,), conn=conn)
            if not _is_pg:
                conn.commit()

    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, _delete)
    return {"ok": True}


@app.post("/admin/users/{user_id}/reset-week")
async def admin_reset_week(user_id: str, current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: _with_conn(lambda c: reset_weekly_limit(c, user_id)))
    return {"ok": True}


@app.get("/sessions/{thread_id}/history")
async def get_session_history(thread_id: str, current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    owns = await loop.run_in_executor(
        None, lambda: session_store.owns_session(thread_id, current_user["id"])
    )
    if not owns:
        raise HTTPException(status_code=404, detail="Session not found")
    config = {"configurable": {"thread_id": thread_id}}
    state_snap = await loop.run_in_executor(
        None, functools.partial(graph.get_state, config)
    )
    if not state_snap or not state_snap.values:
        stored = await loop.run_in_executor(
            None, lambda: session_store.get_history(thread_id, current_user["id"])
        )
        return {"history": stored}
    history = _reconstruct_history(state_snap.values.get("messages", []))
    if not history:
        stored = await loop.run_in_executor(
            None, lambda: session_store.get_history(thread_id, current_user["id"])
        )
        return {"history": stored}
    return {"history": history}


@app.get("/sessions/{thread_id}/scorecard")
async def get_scorecard(thread_id: str, current_user: dict = Depends(get_current_user)):
    loop = asyncio.get_running_loop()
    user, scorecard_text = await asyncio.gather(
        loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, current_user["id"]))),
        loop.run_in_executor(None, lambda: session_store.get_scorecard(thread_id, current_user["id"])),
    )
    if not scorecard_text:
        raise HTTPException(status_code=404, detail="Scorecard not found for this session")
    plan = user.get("plan", "free") if user else "free"
    hire_signal = extract_hire_signal(scorecard_text)
    if plan != "hunter":
        return {"scorecard": None, "hire_signal": hire_signal}
    return {"scorecard": scorecard_text, "hire_signal": hire_signal}


# ── Audio endpoints ───────────────────────────────────────────────────────────

@app.post("/parse-document")
async def parse_document(file: UploadFile):
    data = await file.read()
    name = (file.filename or "").lower()
    text = ""
    try:
        if name.endswith(".pdf"):
            import pdfplumber, io
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif name.endswith(".docx"):
            import docx, io
            doc = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
        else:
            raise HTTPException(status_code=415, detail="Only PDF and DOCX files are supported.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not parse file: {exc}")
    text = text.strip()
    if not text:
        raise HTTPException(status_code=422, detail="No text found in the file.")
    return {"text": text}


@app.post("/stt")
async def stt(audio: UploadFile, lang: str = "en"):
    audio_bytes = await audio.read()
    whisper_lang = "pt" if (lang or "").lower().startswith("pt") else "en"
    transcript = await openai_client.audio.transcriptions.create(
        model="whisper-1",
        file=(audio.filename or "audio.webm", audio_bytes, audio.content_type or "audio/webm"),
        language=whisper_lang,
    )
    return {"text": transcript.text}


_TTS_LANG_INSTRUCTIONS = {
    "pt": "Fale em Português do Brasil com sotaque brasileiro natural e claro.",
    "en": "",
}


def _parse_audio_mime_type(mime_type: str) -> dict:
    bits_per_sample, rate = 16, 24000
    for param in mime_type.split(";"):
        param = param.strip()
        if param.lower().startswith("rate="):
            try:
                rate = int(param.split("=", 1)[1])
            except (ValueError, IndexError):
                pass
        elif param.startswith("audio/L"):
            try:
                bits_per_sample = int(param.split("L", 1)[1])
            except (ValueError, IndexError):
                pass
    return {"bits_per_sample": bits_per_sample, "rate": rate}


def _convert_to_wav(audio_data: bytes, mime_type: str) -> bytes:
    p = _parse_audio_mime_type(mime_type)
    bps, rate, ch = p["bits_per_sample"], p["rate"], 1
    block_align = ch * (bps // 8)
    header = struct.pack(
        "<4sI4s4sIHHIIHH4sI",
        b"RIFF", 36 + len(audio_data), b"WAVE",
        b"fmt ", 16, 1, ch, rate, rate * block_align, block_align, bps,
        b"data", len(audio_data),
    )
    return header + audio_data


_GEMINI_TTS_PROMPT = """\
Read the following transcript based on the audio profile and director's note.

# Audio Profile
An experienced technical recruiter conducting a structured job interview.

# Director's note
Style: Natural. Pace: Measured. Accent: American (Gen).

## Scene:
A professional online job interview. The speaker is an experienced technical recruiter \
conducting a mock interview session to help candidates prepare for real job interviews. \
The interviewer is knowledgeable, organized, and genuinely invested in the candidate's success.

## Sample Context:
Speak with clear pronunciation and a calm, professional cadence. Use a warm but authoritative \
tone — encouraging without being casual. Pause naturally between questions to give the candidate \
space to think. Vary intonation to signal transitions between topics. Sound like a real interviewer, not a narrator.

## Transcript:
{text}"""

_GEMINI_TTS_SCORECARD_PROMPT = """\
Read the following text with energy. Speak at a continuous pace — no long pauses between sentences or sections. Read every word exactly as it appears. Do not add or change anything.

{text}"""


_GEMINI_LANG_CODES = {"en": "en-US", "pt": "pt-BR"}
_GEMINI_VOICES = {
    "Zephyr", "Puck", "Charon", "Kore", "Fenrir", "Leda", "Orus", "Aoede",
    "Callirrhoe", "Autonoe", "Enceladus", "Iapetus", "Umbriel", "Algieba",
    "Despina", "Erinome", "Algenib", "Rasalghul", "Laomedeia", "Achernar",
    "Alnilam", "Schedar", "Gacrux", "Pulcherrima", "Achird", "Zubenelgenubi",
    "Vindemiatrix", "Sadachbia", "Sadaltager", "Sulafat",
}


def _gemini_tts_sync(text: str, lang: str = "en", voice: str | None = None, mode: str = "interview") -> bytes:
    language_code = _GEMINI_LANG_CODES.get(lang, "en-US")
    voice_name = voice if voice in _GEMINI_VOICES else _GEMINI_TTS_VOICE
    prompt_template = _GEMINI_TTS_SCORECARD_PROMPT if mode == "scorecard" else _GEMINI_TTS_PROMPT
    contents = [
        _genai_types.Content(
            role="user",
            parts=[_genai_types.Part.from_text(text=prompt_template.format(text=text))],
        )
    ]
    config = _genai_types.GenerateContentConfig(
        temperature=1,
        response_modalities=["audio"],
        speech_config=_genai_types.SpeechConfig(
            language_code=language_code,
            voice_config=_genai_types.VoiceConfig(
                prebuilt_voice_config=_genai_types.PrebuiltVoiceConfig(
                    voice_name=voice_name,
                )
            )
        ),
    )
    chunks = []
    mime_used = "audio/L16;rate=24000"
    for chunk in _gemini_client.models.generate_content_stream(
        model="gemini-3.1-flash-tts-preview",
        contents=contents,
        config=config,
    ):
        if chunk.parts is None:
            continue
        part = chunk.parts[0]
        if part.inline_data and part.inline_data.data:
            mime_used = part.inline_data.mime_type
            chunks.append(part.inline_data.data)

    raw = b"".join(chunks)
    if mimetypes.guess_extension(mime_used) is None:
        return _convert_to_wav(raw, mime_used)
    return raw


@app.get("/tts-config")
async def tts_config():
    return {"scorecard_voice": _GEMINI_TTS_SCORECARD_VOICE, "default_voice": _GEMINI_TTS_VOICE}


@app.get("/tts")
async def tts(text: str, voice: str = "nova", lang: str = "en", nocache: bool = False, mode: str = "interview"):
    if _TTS_PROVIDER == "gemini":
        effective_gemini_voice = voice if voice in _GEMINI_VOICES else _GEMINI_TTS_VOICE
        ext = "wav"
        media_type = "audio/wav"
    else:
        allowed = {"alloy", "echo", "fable", "onyx", "nova", "shimmer"}
        if voice not in allowed:
            voice = "nova"
        instructions = _TTS_LANG_INSTRUCTIONS.get(lang, "")
        use_mini = bool(instructions)
        tts_model = "gpt-4o-mini-tts" if use_mini else "tts-1"
        ext = "mp3"
        media_type = "audio/mpeg"

    last_exc: Exception | None = None
    for attempt in range(3):
        try:
            if _TTS_PROVIDER == "gemini":
                audio_bytes = await asyncio.to_thread(_gemini_tts_sync, text, lang, voice, mode)
            else:
                kwargs = dict(model=tts_model, voice=voice, input=text)
                if instructions:
                    kwargs["instructions"] = instructions
                response = await openai_client.audio.speech.create(**kwargs)
                audio_bytes = response.content

            return Response(content=audio_bytes, media_type=media_type)
        except Exception as exc:
            last_exc = exc
            if attempt < 2:
                await asyncio.sleep(0.5 * (attempt + 1))

    print(f"[tts] failed after 3 attempts: {last_exc}")
    raise HTTPException(status_code=503, detail="TTS service temporarily unavailable")


# ── History reconstruction from LangGraph state ──────────────────────────────

_OPENING_TEXTS = {
    "Olá, estou pronto para começar a entrevista.",
    "Hi, I'm ready to start the interview.",
}

def _reconstruct_history(messages: list) -> list:
    """Convert LangGraph state messages into display-ready history entries.

    Returns a list of {type, text, voice} dicts — the same format the client
    uses for addBubble / addFeedback.  Internal markers and the scorecard are
    filtered out; report messages (AIMessage followed by [acknowledged...]) are
    tagged as 'feedback' so they render in the green card style.
    """
    history = []
    for i, msg in enumerate(messages):
        if isinstance(msg, HumanMessage):
            content = msg.content
            # Skip the synthetic opening and all internal bracket markers
            if content in _OPENING_TEXTS or content.startswith("["):
                continue
            history.append({"type": "user", "text": content, "voice": "nova"})

        elif isinstance(msg, AIMessage):
            content = msg.content
            if not content or content in ("[no_report]",):
                continue
            if "INTERVIEW SCORECARD" in content:
                continue
            # Detect phase reports: an AIMessage followed by [acknowledged…]
            next_msg = messages[i + 1] if i + 1 < len(messages) else None
            is_report = (
                next_msg is not None
                and isinstance(next_msg, HumanMessage)
                and "[acknowledged" in next_msg.content
            )
            history.append({
                "type": "feedback" if is_report else "ai",
                "text": content,
                "voice": "onyx" if is_report else "nova",
            })
    return history


# ── WebSocket streaming helper ────────────────────────────────────────────────

async def _stream_to_client(ws: WebSocket, messages: list, system_prompt: str,
                             voice: str = "nova", lang: str = "en", max_retries: int = 5,
                             llm=None) -> str:
    from anthropic import APIStatusError
    _llm = llm or model
    full_msg = [SystemMessage(content=[{
        "type": "text",
        "text": system_prompt,
        "cache_control": {"type": "ephemeral"},
    }])] + messages
    for attempt in range(max_retries):
        full_text = ""
        try:
            async for chunk in _llm.astream(full_msg):
                if chunk.content:
                    full_text += chunk.content
                    await ws.send_json({"type": "stream_chunk", "text": chunk.content, "voice": voice, "lang": lang})
            await ws.send_json({"type": "stream_done"})
            return full_text
        except APIStatusError as exc:
            if exc.status_code == 529 and attempt < max_retries - 1:
                wait = 5.0 * (2 ** attempt)
                print(f"[stream retry] Anthropic overloaded (529) — waiting {wait:.0f}s (attempt {attempt + 1}/{max_retries})")
                await asyncio.sleep(wait)
            else:
                print(f"[stream error] {exc}")
                await ws.send_json({"type": "stream_done"})
                return full_text
        except Exception as exc:
            print(f"[stream error] {exc}")
            await ws.send_json({"type": "stream_done"})
            return full_text
    return ""


# ── WebSocket ─────────────────────────────────────────────────────────────────

@app.websocket("/ws")
async def interview_ws(ws: WebSocket):
    print(f"[ws] incoming connection request from {ws.client}")
    await ws.accept()
    print("[ws] connection accepted")
    loop = asyncio.get_running_loop()

    # ── Initial handshake ──────────────────────────────────────────────────
    # Client sends: {"type": "init"|"resume", "token": "...", "job_session_id": "...",
    #                "thread_id": "..." (optional, for resume)}
    thread_id: str = str(uuid.uuid4())
    is_resuming = False
    job_session_id: str = ""
    interview_context: str = ""
    candidate_name: str = "the candidate"
    job_title: str = ""
    company: str = ""

    try:
        init = await asyncio.wait_for(ws.receive_json(), timeout=5.0)
    except Exception:
        await ws.send_json({"type": "error", "text": "No init message received"})
        return

    # Authenticate via token in the init message
    token = init.get("token", "")
    user_info = decode_token_raw(token) if token else None
    if not user_info:
        await ws.send_json({"type": "error", "text": "Unauthorized — please log in again"})
        return

    user_id = user_info["id"]
    job_session_id = init.get("job_session_id", "")

    # Load user status — plan gates scorecard
    user_db = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_user_by_id(c, user_id)))
    user_plan: str     = user_db.get("plan", "free")     if user_db else "free"
    user_language: str = user_db.get("language", "en")   if user_db else "en"

    # Load interview context from the job session
    if job_session_id:
        js = await loop.run_in_executor(
            None, lambda: _with_conn(lambda c: get_job_session(c, job_session_id, user_id))
        )
        if js:
            interview_context = js.get("interview_context", "")
            job_title = js.get("job_title", "")
            company   = js.get("company", "")
            candidate_name = extract_candidate_name(interview_context)
        else:
            print(f"[ws] job_session_id={job_session_id!r} not found for user={user_id!r}")
            await ws.send_json({"type": "error", "text": "Interview session not found — please start a new interview from the onboarding page."})
            return
    else:
        print(f"[ws] no job_session_id for user={user_id!r} — starting without context")
        await ws.send_json({"type": "error", "text": "No interview prepared — please go to the home page and start a new interview."})
        return

    # Check for resume
    if init.get("type") == "resume" and init.get("thread_id"):
        thread_id = init["thread_id"]
        is_resuming = True

    # For new rounds (not resuming), check weekly limit and register the round immediately.
    # This catches "practice again" flows that bypass /prepare.
    if not is_resuming:
        _round_count = await loop.run_in_executor(
            None, lambda: _with_conn(lambda c: count_interviews_this_week(c, user_id))
        )
        _round_limit = PLAN_WEEKLY_LIMITS.get(user_plan, 3)
        if _round_count >= _round_limit:
            await ws.send_json({
                "type": "error",
                "text": f"You've used all {_round_limit} interview rounds this week on your {user_plan.capitalize()} plan.",
                "code": "weekly_limit_reached",
                "plan": user_plan,
                "limit": _round_limit,
            })
            await ws.close(code=1000)
            return
        # Register the round in session_meta immediately so future checks count it.
        _now = int(time.time() * 1000)
        await loop.run_in_executor(
            None,
            lambda: session_store.upsert(
                thread_id, _now, _now, "elevator_pitch", [],
                user_id=user_id, job_session_id=job_session_id or None,
            ),
        )

    config = {"configurable": {"thread_id": thread_id}}
    await ws.send_json({"type": "session", "thread_id": thread_id})

    try:
        if is_resuming:
            state_snap = await loop.run_in_executor(
                None, functools.partial(graph.get_state, config)
            )
            if state_snap and state_snap.values:
                interrupts_list = [
                    intr for task in state_snap.tasks for intr in task.interrupts
                ]
                result = dict(state_snap.values)
                result["__interrupt__"] = interrupts_list
                prev_msg_count       = len(state_snap.values.get("messages", []))
                prev_report_streamed = False
                history = _reconstruct_history(state_snap.values.get("messages", []))
                await ws.send_json({"type": "resumed", "history": history})
                await ws.send_json({"type": "phase", "fase": result.get("fase", "")})
            else:
                is_resuming = False

        if not is_resuming:
            initial_state = make_initial_state(
                interview_context=interview_context,
                candidate_name=candidate_name,
                job_title=job_title,
                company=company,
                language=user_language,
                user_plan=user_plan,
            )
            result = await loop.run_in_executor(
                None, functools.partial(graph.invoke, initial_state, config)
            )
            prev_msg_count       = 1
            prev_report_streamed = False
            skip_advance         = False

        while True:
            messages: list = result.get("messages", [])
            interrupts = result.get("__interrupt__", [])
            fase = result.get("fase", "")

            await ws.send_json({"type": "phase", "fase": fase})

            was_report_streamed  = prev_report_streamed
            prev_report_streamed = False

            new_messages = messages[prev_msg_count:]

            # ── Scorecard gate ─────────────────────────────────────────────
            if fase == "done":
                user_db = await loop.run_in_executor(
                    None, lambda: _with_conn(lambda c: get_user_by_id(c, user_id))
                )
                user_plan = user_db.get("plan", "free") if user_db else "free"

                # Always save the scorecard to DB regardless of plan
                def _is_scorecard(content: str) -> bool:
                    return "INTERVIEW SCORECARD" in content or "SCORECARD DA ENTREVISTA" in content

                scorecard_msg = next(
                    (m.content for m in new_messages
                     if isinstance(m, AIMessage) and _is_scorecard(m.content)),
                    None,
                )
                if scorecard_msg:
                    await loop.run_in_executor(
                        None,
                        lambda: session_store.save_scorecard(thread_id, user_id, scorecard_msg),
                    )

                # Persist full history to session_meta so it survives LangGraph checkpoint loss.
                _full_history = _reconstruct_history(messages)
                await loop.run_in_executor(
                    None,
                    lambda: session_store.save_history(thread_id, user_id, _full_history),
                )

                # Send closing/skip message as feedback (triggers spinner in UI).
                # AIMessages that precede the last HumanMessage were already sent to the
                # client via the interrupt mechanism — exclude them to avoid duplicates.
                _last_human = max(
                    (i for i, m in enumerate(new_messages) if isinstance(m, HumanMessage)),
                    default=-1,
                )
                closing_msgs = [
                    m.content for i, m in enumerate(new_messages)
                    if isinstance(m, AIMessage)
                    and i > _last_human
                    and m.content
                    and m.content not in ("[no_report]",)
                    and not _is_scorecard(m.content)
                ]
                for text in closing_msgs:
                    await ws.send_json({"type": "feedback", "text": text})

                # If nothing to show (candidate_questions was skipped), send skip notice
                if not closing_msgs:
                    _skip_msg = {
                        "pt": "⏭ Fase pulada — seguindo em frente.",
                    }.get(user_language, "⏭ Phase skipped — moving on.")
                    await ws.send_json({"type": "transition", "text": _skip_msg})

                await ws.send_json({"type": "scorecard_ready"})
                await ws.send_json({"type": "done"})
                break

            first_ai_idx = next(
                (i for i, m in enumerate(new_messages) if isinstance(m, AIMessage)), None
            )
            had_transition = any(
                isinstance(m, HumanMessage) and "[acknowledged" in m.content
                for m in new_messages
            )

            for i, msg in enumerate(new_messages):
                if not isinstance(msg, AIMessage):
                    continue
                if msg.content in ("[no_report]", ""):
                    continue
                if i == first_ai_idx:
                    next_msg = new_messages[i + 1] if i + 1 < len(new_messages) else None
                    is_replay = (
                        next_msg is not None
                        and isinstance(next_msg, HumanMessage)
                        and "[acknowledged" not in next_msg.content
                    )
                    if is_replay:
                        continue
                    if was_report_streamed and had_transition:
                        continue
                if _is_scorecard(msg.content):
                    continue  # scorecard served separately via /sessions/{thread_id}/scorecard
                is_scorecard = fase == "done"
                msg_type = "feedback" if is_scorecard else "transition"
                await ws.send_json({"type": msg_type, "text": msg.content})

            if not interrupts:
                if fase == "done" and user_plan == "hunter":
                    await ws.send_json({"type": "scorecard_ready"})
                await ws.send_json({"type": "done"})
                break

            question = interrupts[0].value

            # ── Report streaming ──────────────────────────────────────────
            if question == "[report_ready]":
                phase_key = REPORT_PHASE_MAP.get(fase, "")
                phase_messages = result.get("phase_messages") or []
                lang = result.get("language") or user_language

                # Detect if the candidate skipped this phase
                was_skipped = any(
                    isinstance(m, HumanMessage) and "[Candidate skipped" in m.content
                    for m in phase_messages
                )

                _NEXT_PHASE_LABEL = {
                    "en": {
                        "elevator_pitch_report": "the CAR Project Story",
                        "CAR_report":            "the Technical Questions",
                        "technical_report":      "Leadership & Project Approach",
                        "leadership_report":     "Fit & Motivation",
                        "motivation_report":     "your questions for me",
                    },
                    "pt": {
                        "elevator_pitch_report": "a história de projeto CAR",
                        "CAR_report":            "as Perguntas Técnicas",
                        "technical_report":      "Liderança & Abordagem de Projeto",
                        "leadership_report":     "Fit & Motivação",
                        "motivation_report":     "suas perguntas para mim",
                    },
                }
                _phase_labels = _NEXT_PHASE_LABEL.get(lang, _NEXT_PHASE_LABEL["en"])
                next_label = _phase_labels.get(fase, "the next phase" if lang != "pt" else "a próxima fase")

                _READY_MSG = {
                    "pt": f"Tome um momento para revisar esse feedback. Quando estiver pronto para seguir para {next_label}, é só me enviar uma mensagem.",
                    "en": f"Take a moment to review that feedback. When you're ready to move on to {next_label}, just send me a message.",
                }

                if was_skipped:
                    full_text    = "[no_report]"
                    skip_advance = True
                    _skip_msg = {
                        "pt": "⏭ Fase pulada — seguindo em frente.",
                    }.get(lang, "⏭ Phase skipped — moving on.")
                    await ws.send_json({"type": "ai", "text": _skip_msg})
                elif phase_key:
                    # Build dynamic report system prompt with per-session context
                    ctx = result.get("interview_context") or interview_context
                    report_system = _build_report_system(ctx, phase_key, lang)

                    # Format phase transcript as a single human message so the
                    # model reads it as a document to evaluate, not a conversation
                    # to continue. Passing raw messages makes the model reply as
                    # Alex (asking another question) instead of as the Judge.
                    _OPENING_TEXTS_SET = {
                        "I'm ready to start this phase.",
                        "Olá, estou pronto para começar a entrevista.",
                        "Hi, I'm ready to start the interview.",
                    }
                    transcript_lines = []
                    for m in phase_messages:
                        if isinstance(m, HumanMessage):
                            if m.content in _OPENING_TEXTS_SET or m.content.startswith("["):
                                continue
                            transcript_lines.append(f"CANDIDATE: {m.content}")
                        elif isinstance(m, AIMessage):
                            if not m.content or m.content.startswith("["):
                                continue
                            transcript_lines.append(f"INTERVIEWER: {m.content}")
                    transcript_text = "\n\n".join(transcript_lines)
                    report_messages = [HumanMessage(content=f"## Phase Transcript\n\n{transcript_text}\n\n---\n\nNow generate the structured phase feedback report.")]

                    full_text = await _stream_to_client(
                        ws, report_messages, report_system, voice="onyx", lang=lang,
                        llm=get_model(user_plan),
                    )
                    await ws.send_json({
                        "type": "ai",
                        "text": _READY_MSG.get(lang, _READY_MSG["en"]),
                    })
                else:
                    full_text = ""
                prev_msg_count       = len(messages)
                prev_report_streamed = bool(full_text) and not was_skipped
                result = await loop.run_in_executor(
                    None,
                    functools.partial(
                        graph.invoke, Command(resume=full_text or "[no_report]"), config
                    ),
                )
                continue

            if had_transition and not skip_advance:
                await ws.send_json({"type": "await_input"})
                while True:
                    data = await ws.receive_json()
                    if data.get("type") == "ping":
                        await ws.send_json({"type": "pong"})
                        continue
                    ready_text = data.get("text", "").strip()
                    if ready_text:
                        break
                await ws.send_json({"type": "user", "text": ready_text})
                await ws.send_json({"type": "ai", "text": question})
            else:
                skip_advance = False
                await ws.send_json({"type": "ai", "text": question})

            while True:
                data = await ws.receive_json()
                if data.get("type") == "ping":
                    await ws.send_json({"type": "pong"})
                    continue
                user_text = data.get("text", "").strip()
                if not user_text:
                    continue
                break

            await ws.send_json({"type": "user", "text": user_text})
            prev_msg_count = len(messages)

            try:
                result = await loop.run_in_executor(
                    None,
                    functools.partial(graph.invoke, Command(resume=user_text), config),
                )
            except Exception as exc:
                import traceback
                traceback.print_exc()
                await ws.send_json({
                    "type": "error",
                    "text": f"[Server error — please reload and try again]\n{exc}",
                })
                break

    except WebSocketDisconnect:
        pass
    except Exception as exc:
        import traceback
        traceback.print_exc()
        try:
            await ws.send_json({
                "type": "error",
                "text": f"[Connection error — please reload]\n{exc}",
            })
        except Exception:
            pass


# ── Contact endpoints ─────────────────────────────────────────────────────────

class ContactRequest(BaseModel):
    name:    str
    email:   str
    subject: str
    body:    str


@app.post("/contact")
async def submit_contact(body: ContactRequest):
    if not body.name.strip():
        raise HTTPException(status_code=422, detail="Name is required")
    if not body.email.strip() or "@" not in body.email:
        raise HTTPException(status_code=422, detail="Valid email is required")
    if not body.subject.strip():
        raise HTTPException(status_code=422, detail="Subject is required")
    if len(body.body.strip()) < 10:
        raise HTTPException(status_code=422, detail="Message is too short")

    loop = asyncio.get_running_loop()
    msg_id = await loop.run_in_executor(
        None,
        lambda: _with_conn(lambda c: create_contact_message(c, body.name, body.email, body.subject, body.body)),
    )
    try:
        await loop.run_in_executor(
            None,
            lambda: send_contact_notification(body.name, body.email, body.subject, body.body),
        )
    except Exception as exc:
        print(f"[contact] Failed to send notification: {exc}")
    return {"ok": True, "id": msg_id}


@app.get("/admin/contact")
async def admin_list_contact(current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    loop = asyncio.get_running_loop()
    messages = await loop.run_in_executor(None, lambda: _with_conn(lambda c: list_contact_messages(c)))
    return {"messages": messages}


class ContactReplyRequest(BaseModel):
    reply_text: str


@app.post("/admin/contact/{msg_id}/reply")
async def admin_reply_contact(msg_id: str, body: ContactReplyRequest,
                              current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    if not body.reply_text.strip():
        raise HTTPException(status_code=422, detail="Reply text is required")

    loop = asyncio.get_running_loop()
    msg = await loop.run_in_executor(None, lambda: _with_conn(lambda c: get_contact_message(c, msg_id)))
    if not msg:
        raise HTTPException(status_code=404, detail="Message not found")

    await loop.run_in_executor(
        None, lambda: _with_conn(lambda c: save_contact_reply(c, msg_id, body.reply_text))
    )
    try:
        await loop.run_in_executor(
            None,
            lambda: send_contact_reply(msg["email"], msg["name"], msg["subject"], body.reply_text),
        )
    except Exception as exc:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=503, detail=f"Reply saved but email failed: {exc}")
    return {"ok": True}


@app.delete("/admin/contact/{msg_id}")
async def admin_delete_contact(msg_id: str, current_user: dict = Depends(get_current_user)):
    if current_user.get("email") != ADMIN_EMAIL:
        raise HTTPException(status_code=403, detail="Admin only")
    loop = asyncio.get_running_loop()
    deleted = await loop.run_in_executor(None, lambda: _with_conn(lambda c: delete_contact_message(c, msg_id)))
    if not deleted:
        raise HTTPException(status_code=404, detail="Message not found")
    return {"ok": True}


@app.get("/scorecard/{thread_id}")
async def scorecard_page(thread_id: str):
    return FileResponse(Path(__file__).parent / "static" / "scorecard.html")


@app.get("/")
async def landing_root(request: Request):
    lang = _preferred_landing_language(request)
    response = RedirectResponse(url=f"/{lang}/", status_code=302)
    response.headers["Vary"] = "Accept-Language, Cookie"
    return response


@app.get("/en")
@app.get("/en/")
async def landing_en():
    return FileResponse(Path(__file__).parent / "static" / "index.html")


@app.get("/pt")
@app.get("/pt/")
async def landing_pt():
    return FileResponse(Path(__file__).parent / "static" / "index.html")


# Serve the frontend — mount last so API routes are registered first
STATIC_DIR = Path(__file__).parent / "static"
app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
